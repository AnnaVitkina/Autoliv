"""LTL Excel writers — Matrix/HERO with bracket costs + Conditions."""

from __future__ import annotations

from pathlib import Path

import pandas as pd

from flows.ftl import paths as project_paths
from flows.ltl.config import (
    ATTR_DUPLICATE_LANE,
    ATTR_FORWARDER_REPLACED,
    CURRENCY_COL,
    DETAIL_COLUMNS,
    DETAIL_HEADER_ALIASES,
    SHEET_CONDITIONS,
)
from flows.ltl.cost_utils import block_data_keys, bracket_key, currency_key
from flows.ftl.schema import (
    ATTR_CHANGED_CELLS,
    ATTR_NEW_LANE,
    ATTR_REMOVED_LANE,
    ATTR_VALID_FROM_RATES,
)


def write_matrix_sheet(ws, matrix: pd.DataFrame) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    cost_blocks: list[dict] = list(matrix.attrs.get("cost_blocks", []))
    duplicate_flags: list[bool] = list(matrix.attrs.get(ATTR_DUPLICATE_LANE, []))
    replaced_flags: list[bool] = list(matrix.attrs.get(ATTR_FORWARDER_REPLACED, []))
    changed_rows: list[set] = list(matrix.attrs.get(ATTR_CHANGED_CELLS, []))
    new_flags: list[bool] = list(matrix.attrs.get(ATTR_NEW_LANE, []))
    removed_flags: list[bool] = list(matrix.attrs.get(ATTR_REMOVED_LANE, []))
    valid_from_rates_flags: list[bool] = list(
        matrix.attrs.get(ATTR_VALID_FROM_RATES, [])
    )

    header_rows = 3
    detail_headers = [
        DETAIL_HEADER_ALIASES.get(c, c) for c in DETAIL_COLUMNS
    ]
    detail_count = len(DETAIL_COLUMNS)

    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    detail_fill = PatternFill("solid", fgColor="F2F2F2")
    cost_fill = PatternFill("solid", fgColor="D9D2E9")
    grouped_fill = PatternFill("solid", fgColor="CFE2F3")
    dup_fill = PatternFill("solid", fgColor="FFC7CE")
    common_fill = PatternFill("solid", fgColor="D9D9D9")
    change_fill = PatternFill("solid", fgColor="C6EFCE")
    new_fill = PatternFill("solid", fgColor="FFEB9C")
    removed_fill = PatternFill("solid", fgColor="F4B183")
    valid_from_fill = PatternFill("solid", fgColor="E2D5F1")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    dup_font = Font(color="9C0006")
    conditioned_font = Font(underline="single", color="595959")
    change_font = Font(color="006100")
    new_font = Font(color="9C5700")
    removed_font = Font(color="833C0C")
    valid_from_font = Font(color="7030A0")

    for col_idx, title in enumerate(detail_headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.fill = detail_fill
        cell.alignment = center
        cell.font = Font(bold=True)
        cell.border = border
        ws.merge_cells(
            start_row=1, start_column=col_idx, end_row=header_rows, end_column=col_idx
        )
        for row in range(1, header_rows + 1):
            ws.cell(row=row, column=col_idx).border = border
            ws.cell(row=row, column=col_idx).fill = detail_fill

    col_idx = detail_count + 1
    for block in cost_blocks:
        if block.get("kind") == "grouped":
            children = list(block.get("children", []))
            child_spans = [1 + len(ch.get("brackets", [])) for ch in children]
            total_span = sum(child_spans) if child_spans else 1
            left = col_idx
            right = col_idx + total_span - 1

            name_cell = ws.cell(row=1, column=left, value=block["name"])
            name_cell.fill = grouped_fill
            name_cell.alignment = center
            name_cell.font = Font(bold=True)
            if total_span > 1:
                ws.merge_cells(
                    start_row=1, start_column=left, end_row=1, end_column=right
                )

            child_left = left
            for child, span in zip(children, child_spans):
                child_right = child_left + span - 1
                child_cell = ws.cell(row=2, column=child_left, value=child["name"])
                child_cell.fill = grouped_fill
                child_cell.alignment = center
                child_cell.font = Font(bold=True)
                if span > 1:
                    ws.merge_cells(
                        start_row=2,
                        start_column=child_left,
                        end_row=2,
                        end_column=child_right,
                    )
                sub_headers = [CURRENCY_COL] + list(child.get("brackets", []))
                for offset, sub in enumerate(sub_headers):
                    cell = ws.cell(row=3, column=child_left + offset, value=sub)
                    cell.fill = grouped_fill
                    cell.alignment = center
                    cell.font = Font(bold=True)
                child_left = child_right + 1

            for row in range(1, header_rows + 1):
                for c in range(left, right + 1):
                    ws.cell(row=row, column=c).border = border
                    ws.cell(row=row, column=c).fill = grouped_fill

            block["_data_keys"] = block_data_keys(block)
            col_idx = right + 1
            continue

        name = block["name"]
        brackets = list(block.get("brackets", []))
        span = 1 + len(brackets)
        sub_headers = [CURRENCY_COL] + brackets
        data_keys = block_data_keys(block)
        left = col_idx
        right = col_idx + span - 1

        name_cell = ws.cell(row=1, column=left, value=name)
        name_cell.fill = cost_fill
        name_cell.alignment = center
        name_cell.font = Font(bold=True)
        if span > 1:
            ws.merge_cells(start_row=1, start_column=left, end_row=1, end_column=right)

        rate_by = block.get("rate_by", "")
        rate_cell = ws.cell(row=2, column=left, value=rate_by)
        rate_cell.fill = cost_fill
        rate_cell.alignment = center
        if span > 1:
            ws.merge_cells(start_row=2, start_column=left, end_row=2, end_column=right)

        for offset, sub in enumerate(sub_headers):
            cell = ws.cell(row=3, column=left + offset, value=sub)
            cell.fill = cost_fill
            cell.alignment = center
            cell.font = Font(bold=True)

        for row in range(1, header_rows + 1):
            for c in range(left, right + 1):
                ws.cell(row=row, column=c).border = border
                ws.cell(row=row, column=c).fill = cost_fill

        block["_data_keys"] = data_keys
        col_idx = right + 1

    for r_idx, (_, row) in enumerate(matrix.iterrows()):
        excel_row = header_rows + 1 + r_idx
        is_duplicate = r_idx < len(duplicate_flags) and duplicate_flags[r_idx]
        is_replaced = r_idx < len(replaced_flags) and replaced_flags[r_idx]
        is_new = r_idx < len(new_flags) and new_flags[r_idx]
        is_removed = r_idx < len(removed_flags) and removed_flags[r_idx]
        changed = changed_rows[r_idx] if r_idx < len(changed_rows) else set()

        for c_idx, col_name in enumerate(DETAIL_COLUMNS, start=1):
            value = row[col_name]
            cell = ws.cell(
                row=excel_row,
                column=c_idx,
                value="" if pd.isna(value) else value,
            )
            cell.border = border
            if is_removed:
                cell.fill = removed_fill
                cell.font = removed_font
            elif (
                col_name == "Valid from"
                and r_idx < len(valid_from_rates_flags)
                and valid_from_rates_flags[r_idx]
            ):
                cell.fill = valid_from_fill
                cell.font = valid_from_font
            elif (not is_new) and col_name in changed:
                cell.fill = change_fill
                cell.font = (
                    Font(underline="single", color="006100")
                    if is_replaced and col_name == "Forwarder"
                    else change_font
                )
            elif is_new:
                cell.fill = new_fill
                cell.font = (
                    Font(underline="single", color="9C5700")
                    if is_replaced and col_name == "Forwarder"
                    else new_font
                )
            elif is_duplicate:
                cell.fill = dup_fill
                cell.font = (
                    Font(underline="single", color="9C0006")
                    if is_replaced and col_name == "Forwarder"
                    else dup_font
                )
            elif is_replaced and col_name == "Forwarder":
                cell.fill = common_fill
                cell.font = conditioned_font

        col_idx = detail_count + 1
        for block in cost_blocks:
            for key in block.get("_data_keys", []):
                value = row[key] if key in matrix.columns else ""
                cell = ws.cell(
                    row=excel_row,
                    column=col_idx,
                    value="" if pd.isna(value) else value,
                )
                cell.border = border
                if is_removed:
                    cell.fill = removed_fill
                    cell.font = removed_font
                elif (not is_new) and key in changed:
                    cell.fill = change_fill
                    cell.font = change_font
                elif is_new:
                    cell.fill = new_fill
                    cell.font = new_font
                elif is_duplicate:
                    cell.fill = dup_fill
                    cell.font = dup_font
                col_idx += 1

    total_cols = col_idx - 1
    for c_idx in range(1, detail_count + 1):
        ws.column_dimensions[get_column_letter(c_idx)].width = 14
    for c_idx in range(detail_count + 1, max(total_cols, detail_count) + 1):
        ws.column_dimensions[get_column_letter(c_idx)].width = 10

    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 18
    if total_cols >= 1:
        ws.auto_filter.ref = (
            f"A{header_rows}:{get_column_letter(max(total_cols, 1))}{header_rows}"
        )
    ws.freeze_panes = f"A{header_rows + 1}"


def write_conditions_sheet(ws, rules_df: pd.DataFrame) -> None:
    from openpyxl.styles import Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="D9D2E9")
    section_fill = PatternFill("solid", fgColor="F2F2F2")
    grey_fill = PatternFill("solid", fgColor="D9D9D9")
    conditioned_font = Font(underline="single", color="595959")
    bold = Font(bold=True)

    applied_df = rules_df.attrs.get("applied_replacements")
    if applied_df is None:
        applied_df = pd.DataFrame(
            columns=["Field", "Original value", "Common value", "Condition", "Rows"]
        )

    row = 1
    ws.cell(row=row, column=1, value="Common rating rules – Forwarder").font = Font(
        bold=True, size=12
    )
    row = 3
    rule_headers = ["Field", "Common value", "Condition"]
    for c_idx, title in enumerate(rule_headers, start=1):
        cell = ws.cell(row=row, column=c_idx, value=title)
        cell.font = bold
        cell.fill = header_fill
        cell.border = border
    row += 1
    for _, rec in rules_df.iterrows():
        for c_idx, key in enumerate(rule_headers, start=1):
            cell = ws.cell(row=row, column=c_idx, value=rec[key])
            cell.border = border
            if key == "Common value":
                cell.fill = grey_fill
                cell.font = conditioned_font
        row += 1

    row += 2
    ws.cell(row=row, column=1, value="Replacements applied in this file").font = Font(
        bold=True, size=12
    )
    row += 2
    applied_headers = ["Field", "Original value", "Common value", "Condition", "Rows"]
    for c_idx, title in enumerate(applied_headers, start=1):
        cell = ws.cell(row=row, column=c_idx, value=title)
        cell.font = bold
        cell.fill = section_fill
        cell.border = border
    row += 1
    if len(applied_df) == 0:
        ws.cell(row=row, column=1, value="(no Forwarder replacements in filtered lanes)")
    else:
        for _, rec in applied_df.iterrows():
            for c_idx, key in enumerate(applied_headers, start=1):
                cell = ws.cell(row=row, column=c_idx, value=rec[key])
                cell.border = border
                if key == "Common value":
                    cell.fill = grey_fill
                    cell.font = conditioned_font
            row += 1

    widths = [16, 40, 55, 12, 8]
    for c_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c_idx)].width = width


def save_workbook(
    sheets: dict[str, pd.DataFrame], rates_stem: str
) -> Path:
    from openpyxl import Workbook

    project_paths.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out_path = project_paths.OUTPUT_DIR / f"{rates_stem}_matrix.xlsx"

    wb = Workbook()
    default = wb.active
    wb.remove(default)

    for sheet_name, matrix in sheets.items():
        ws = wb.create_sheet(sheet_name)
        if sheet_name == SHEET_CONDITIONS:
            write_conditions_sheet(ws, matrix)
        else:
            write_matrix_sheet(ws, matrix)

    wb.save(out_path)
    return out_path
