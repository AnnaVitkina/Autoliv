"""Excel writers for matrix sheets and Conditions tab."""

from __future__ import annotations

from pathlib import Path

import pandas as pd

from flows.ftl.cost_utils import currency_key, flat_key
from flows.ftl import paths as project_paths
from flows.ftl.schema import (
    ATTR_CHANGED_CELLS,
    ATTR_COLLECTION_COMBINED,
    ATTR_COMPANY_CODE_CONDITIONS,
    ATTR_DEST_COMBINED,
    ATTR_DUPLICATE_LANE,
    ATTR_FORWARDER_REPLACED,
    ATTR_NEW_LANE,
    ATTR_ORIGIN_COMBINED,
    ATTR_PAYING_PLANT_COMBINED,
    ATTR_REMOVED_LANE,
    ATTR_VALID_FROM_RATES,
    CURRENCY_COL,
    DETAIL_COLUMNS,
    FLAT_COL,
    RATE_BY_LABEL,
    SHEET_CONDITIONS,
)


def _empty_generated_conditions() -> pd.DataFrame:
    return pd.DataFrame(
        columns=[
            "Field",
            "Original value",
            "Attribute",
            "Operator",
            "Value",
            "Scope",
            "Rows",
        ]
    )


def write_matrix_sheet(ws, matrix: pd.DataFrame) -> None:
    """Write one sheet with multi-level cost headers."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    cost_names: list[str] = list(matrix.attrs.get("cost_names", []))
    if not cost_names:
        for col in matrix.columns:
            if isinstance(col, str) and col.endswith(f"||{CURRENCY_COL}"):
                cost_names.append(col.split("||", 1)[0])

    replaced_flags: list[bool] = list(matrix.attrs.get(ATTR_FORWARDER_REPLACED, []))
    origin_flags: list[bool] = list(matrix.attrs.get(ATTR_ORIGIN_COMBINED, []))
    dest_flags: list[bool] = list(matrix.attrs.get(ATTR_DEST_COMBINED, []))
    collection_flags: list[bool] = list(
        matrix.attrs.get(ATTR_COLLECTION_COMBINED, [])
    )
    paying_flags: list[bool] = list(
        matrix.attrs.get(ATTR_PAYING_PLANT_COMBINED, [])
    )
    changed_rows: list[set] = list(matrix.attrs.get(ATTR_CHANGED_CELLS, []))
    new_flags: list[bool] = list(matrix.attrs.get(ATTR_NEW_LANE, []))
    removed_flags: list[bool] = list(matrix.attrs.get(ATTR_REMOVED_LANE, []))
    duplicate_flags: list[bool] = list(matrix.attrs.get(ATTR_DUPLICATE_LANE, []))
    valid_from_rates_flags: list[bool] = list(
        matrix.attrs.get(ATTR_VALID_FROM_RATES, [])
    )

    header_rows = 3
    detail_headers = [
        "Paying_Plant" if c == "Paying_Plant_2" else c for c in DETAIL_COLUMNS
    ]
    detail_count = len(DETAIL_COLUMNS)

    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    detail_fill = PatternFill("solid", fgColor="F2F2F2")
    cost_fill = PatternFill("solid", fgColor="D9D2E9")
    common_fill = PatternFill("solid", fgColor="D9D9D9")
    change_fill = PatternFill("solid", fgColor="C6EFCE")
    new_fill = PatternFill("solid", fgColor="FFEB9C")
    dup_fill = PatternFill("solid", fgColor="FFC7CE")
    removed_fill = PatternFill("solid", fgColor="F4B183")
    valid_from_fill = PatternFill("solid", fgColor="E2D5F1")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    conditioned_font = Font(underline="single", color="595959")
    change_font = Font(color="006100")
    new_font = Font(color="9C5700")
    dup_font = Font(color="9C0006")
    removed_font = Font(color="833C0C")
    valid_from_font = Font(color="7030A0")

    for col_idx, title in enumerate(detail_headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.fill = detail_fill
        cell.alignment = center
        cell.font = Font(bold=True)
        cell.border = border
        ws.merge_cells(
            start_row=1, start_column=col_idx, end_row=header_rows, end_column=col_idx
        )
        for row in range(1, header_rows + 1):
            ws.cell(row=row, column=col_idx).border = border
            ws.cell(row=row, column=col_idx).fill = detail_fill

    col_idx = detail_count + 1
    for cost_name in cost_names:
        left = col_idx
        right = col_idx + 1

        name_cell = ws.cell(row=1, column=left, value=cost_name)
        name_cell.fill = cost_fill
        name_cell.alignment = center
        name_cell.font = Font(bold=True)
        ws.merge_cells(start_row=1, start_column=left, end_row=1, end_column=right)

        rate_cell = ws.cell(row=2, column=left, value=RATE_BY_LABEL)
        rate_cell.fill = cost_fill
        rate_cell.alignment = center
        ws.merge_cells(start_row=2, start_column=left, end_row=2, end_column=right)

        for offset, sub in enumerate((CURRENCY_COL, FLAT_COL)):
            cell = ws.cell(row=3, column=left + offset, value=sub)
            cell.fill = cost_fill
            cell.alignment = center
            cell.font = Font(bold=True)

        for row in range(1, header_rows + 1):
            for c in (left, right):
                ws.cell(row=row, column=c).border = border
                ws.cell(row=row, column=c).fill = cost_fill

        col_idx += 2

    for r_idx, (_, row) in enumerate(matrix.iterrows()):
        excel_row = header_rows + 1 + r_idx
        changed = changed_rows[r_idx] if r_idx < len(changed_rows) else set()
        is_new = r_idx < len(new_flags) and new_flags[r_idx]
        is_removed = r_idx < len(removed_flags) and removed_flags[r_idx]
        is_duplicate = r_idx < len(duplicate_flags) and duplicate_flags[r_idx]

        for c_idx, col_name in enumerate(DETAIL_COLUMNS, start=1):
            value = row[col_name]
            cell = ws.cell(
                row=excel_row,
                column=c_idx,
                value="" if pd.isna(value) else value,
            )
            cell.border = border

            is_conditioned = (
                (
                    col_name == "Forwarder"
                    and r_idx < len(replaced_flags)
                    and replaced_flags[r_idx]
                )
                or (
                    col_name == "OriginNo"
                    and r_idx < len(origin_flags)
                    and origin_flags[r_idx]
                )
                or (
                    col_name == "DestinationNo"
                    and r_idx < len(dest_flags)
                    and dest_flags[r_idx]
                )
                or (
                    col_name == "Collection day"
                    and r_idx < len(collection_flags)
                    and collection_flags[r_idx]
                )
                or (
                    col_name == "Paying_Plant"
                    and r_idx < len(paying_flags)
                    and paying_flags[r_idx]
                )
            )
            is_rates_valid_from = (
                col_name == "Valid from"
                and r_idx < len(valid_from_rates_flags)
                and valid_from_rates_flags[r_idx]
            )
            if is_removed:
                cell.fill = removed_fill
                cell.font = removed_font
            elif is_rates_valid_from:
                cell.fill = valid_from_fill
                cell.font = valid_from_font
            elif (not is_new) and col_name in changed:
                cell.fill = change_fill
                cell.font = (
                    Font(underline="single", color="006100")
                    if is_conditioned
                    else change_font
                )
            elif is_new:
                cell.fill = new_fill
                cell.font = (
                    Font(underline="single", color="9C5700")
                    if is_conditioned
                    else new_font
                )
            elif is_duplicate:
                cell.fill = dup_fill
                cell.font = (
                    Font(underline="single", color="9C0006")
                    if is_conditioned
                    else dup_font
                )
            elif is_conditioned:
                cell.fill = common_fill
                cell.font = conditioned_font

        col_idx = detail_count + 1
        for cost_name in cost_names:
            for key in (currency_key(cost_name), flat_key(cost_name)):
                value = row[key] if key in matrix.columns else ""
                cell = ws.cell(
                    row=excel_row,
                    column=col_idx,
                    value="" if pd.isna(value) else value,
                )
                cell.border = border
                if is_removed:
                    cell.fill = removed_fill
                    cell.font = removed_font
                elif (not is_new) and key in changed:
                    cell.fill = change_fill
                    cell.font = change_font
                elif is_new:
                    cell.fill = new_fill
                    cell.font = new_font
                elif is_duplicate:
                    cell.fill = dup_fill
                    cell.font = dup_font
                col_idx += 1

    total_cols = detail_count + len(cost_names) * 2
    for c_idx in range(1, detail_count + 1):
        ws.column_dimensions[get_column_letter(c_idx)].width = 14
    for c_idx in range(detail_count + 1, total_cols + 1):
        ws.column_dimensions[get_column_letter(c_idx)].width = 12

    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 18
    if total_cols >= 1:
        ws.auto_filter.ref = (
            f"A{header_rows}:{get_column_letter(max(total_cols, 1))}{header_rows}"
        )
    ws.freeze_panes = f"A{header_rows + 1}"


def write_conditions_sheet(ws, rules_df: pd.DataFrame) -> None:
    """Write Forwarder common-rating rules and Origin/Destination conditions."""
    from openpyxl.styles import Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="D9D2E9")
    section_fill = PatternFill("solid", fgColor="F2F2F2")
    grey_fill = PatternFill("solid", fgColor="D9D9D9")
    conditioned_font = Font(underline="single", color="595959")
    bold = Font(bold=True)

    applied_df = rules_df.attrs.get("applied_replacements")
    if applied_df is None:
        applied_df = pd.DataFrame(
            columns=["Field", "Original value", "Common value", "Condition", "Rows"]
        )

    company_df = rules_df.attrs.get(ATTR_COMPANY_CODE_CONDITIONS)
    if company_df is None:
        company_df = _empty_generated_conditions()

    def style_common_cell(cell) -> None:
        cell.fill = grey_fill
        cell.font = conditioned_font

    row = 1
    ws.cell(row=row, column=1, value="Common rating rules – Forwarder").font = Font(
        bold=True, size=12
    )
    row = 3
    rule_headers = ["Field", "Common value", "Condition"]
    for c_idx, title in enumerate(rule_headers, start=1):
        cell = ws.cell(row=row, column=c_idx, value=title)
        cell.font = bold
        cell.fill = header_fill
        cell.border = border
    row += 1
    for _, rec in rules_df.iterrows():
        for c_idx, key in enumerate(rule_headers, start=1):
            cell = ws.cell(row=row, column=c_idx, value=rec[key])
            cell.border = border
            if key == "Common value":
                style_common_cell(cell)
        row += 1

    row += 2
    ws.cell(row=row, column=1, value="Replacements applied in this file").font = Font(
        bold=True, size=12
    )
    row += 2
    applied_headers = ["Field", "Original value", "Common value", "Condition", "Rows"]
    for c_idx, title in enumerate(applied_headers, start=1):
        cell = ws.cell(row=row, column=c_idx, value=title)
        cell.font = bold
        cell.fill = section_fill
        cell.border = border
    row += 1

    if len(applied_df) == 0:
        ws.cell(row=row, column=1, value="(no Forwarder replacements in filtered lanes)")
        row += 1
    else:
        for _, rec in applied_df.iterrows():
            for c_idx, key in enumerate(applied_headers, start=1):
                cell = ws.cell(row=row, column=c_idx, value=rec[key])
                cell.border = border
                if key == "Common value":
                    style_common_cell(cell)
            row += 1

    row += 2
    ws.cell(
        row=row,
        column=1,
        value="Generated conditions (4 columns: Attribute | Operator | Value | Scope)",
    ).font = Font(bold=True, size=12)
    row += 2
    company_headers = [
        "Field",
        "Original value",
        "Attribute",
        "Operator",
        "Value",
        "Scope",
        "Rows",
    ]
    for c_idx, title in enumerate(company_headers, start=1):
        cell = ws.cell(row=row, column=c_idx, value=title)
        cell.font = bold
        cell.fill = header_fill
        cell.border = border
    row += 1

    if len(company_df) == 0:
        ws.cell(
            row=row,
            column=1,
            value="(no generated Origin / Destination / Collection day / Paying Plant conditions)",
        )
    else:
        for _, rec in company_df.iterrows():
            for c_idx, key in enumerate(company_headers, start=1):
                cell = ws.cell(row=row, column=c_idx, value=rec[key])
                cell.border = border
                if key in {"Original value", "Value"}:
                    style_common_cell(cell)
            row += 1

    widths = [16, 40, 22, 12, 45, 14, 8]
    for c_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c_idx)].width = width


def save_workbook(sheets: dict[str, pd.DataFrame], rates_file: Path) -> Path:
    from openpyxl import Workbook

    project_paths.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out_path = project_paths.OUTPUT_DIR / f"{rates_file.stem}_matrix.xlsx"

    wb = Workbook()
    default = wb.active
    wb.remove(default)

    for sheet_name, matrix in sheets.items():
        ws = wb.create_sheet(sheet_name)
        if sheet_name == SHEET_CONDITIONS:
            write_conditions_sheet(ws, matrix)
        else:
            write_matrix_sheet(ws, matrix)

    wb.save(out_path)
    return out_path
