"""Previous-RA load, content-aware reorder, and change highlighting."""

from __future__ import annotations

from collections import defaultdict
from pathlib import Path

import pandas as pd

from flows.ftl.cost_utils import cost_names_match, currency_key, flat_key
from flows.ftl.schema import (
    ATTR_CHANGED_CELLS,
    ATTR_DUPLICATE_LANE,
    ATTR_NEW_LANE,
    ATTR_RA_MATCH,
    CURRENCY_COL,
    DETAIL_COLUMNS,
    FLAT_COL,
    LANE_MATCH_COLS,
    PER_ROW_ATTRS,
    RA_DETAIL_HEADER_MAP,
    RATE_CARD_SHEET,
    SHEET_CONDITIONS,
    SHEET_HERO,
    SHEET_MATRIX,
    SHEET_OTHER,
)
from flows.ftl.values import cmp_value


def transport_flat_from_row(
    row: pd.Series, matrix: pd.DataFrame, transport_cost_name: str
) -> object:
    key = flat_key(transport_cost_name)
    if key in matrix.columns:
        return row[key]
    return ""


def transport_flat_from_ra(ra_lane: dict, transport_cost_name: str) -> object:
    for ra_name, vals in ra_lane.get("costs", {}).items():
        if cost_names_match(transport_cost_name, ra_name):
            return vals.get(FLAT_COL, "")
    return ""


def lane_match_score(
    row: pd.Series,
    ra_lane: dict,
    matrix: pd.DataFrame,
    transport_cost_name: str,
) -> int:
    """
    Score how well a matrix row matches a previous-RA lane that shares RateCode.
    Higher is better. Used to pair Single/Round (and similar) twins correctly.
    """
    score = 0
    details = ra_lane.get("details", {})
    for col in LANE_MATCH_COLS:
        ours = cmp_value(row[col] if col in matrix.columns else "")
        theirs = cmp_value(details.get(col, ""))
        if ours == "" and theirs == "":
            continue
        if ours == theirs:
            if col in ("Singletrip_Roundtrip", "Collection day", "Service"):
                score += 10
            else:
                score += 1
        elif col in ("Singletrip_Roundtrip", "Collection day") and ours and theirs:
            score -= 20

    our_flat = cmp_value(transport_flat_from_row(row, matrix, transport_cost_name))
    ra_flat = cmp_value(transport_flat_from_ra(ra_lane, transport_cost_name))
    if our_flat != "" and ra_flat != "":
        if our_flat == ra_flat:
            score += 15
        else:
            score -= 5
    return score


def pick_best_lane_index(
    candidate_indices: list[int],
    ra_lane: dict,
    matrix: pd.DataFrame,
    transport_cost_name: str,
) -> int | None:
    """Pick the unused matrix row that best matches this RA lane."""
    if not candidate_indices:
        return None
    best_idx = candidate_indices[0]
    best_score = lane_match_score(
        matrix.iloc[best_idx], ra_lane, matrix, transport_cost_name
    )
    for idx in candidate_indices[1:]:
        score = lane_match_score(
            matrix.iloc[idx], ra_lane, matrix, transport_cost_name
        )
        if score > best_score:
            best_idx = idx
            best_score = score
    return best_idx


def load_previous_ra_rate_card(path: Path) -> list[dict]:
    """
    Read previous RA 'Rate card' tab into ordered lane dicts:
      {RateCode, details: {our_col: value}, costs: {ra_cost_name: {Currency, Flat}}}
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    if RATE_CARD_SHEET not in wb.sheetnames:
        raise KeyError(
            f"Sheet '{RATE_CARD_SHEET}' not found in {path.name}. "
            f"Available: {wb.sheetnames}"
        )
    ws = wb[RATE_CARD_SHEET]

    header_row = None
    for r in range(1, min(20, ws.max_row + 1)):
        vals = [str(ws.cell(r, c).value or "").strip() for c in range(1, 30)]
        if "Lane #" in vals and "RateCode" in vals:
            header_row = r
            break
    if header_row is None:
        raise ValueError(f"Could not find Rate card header row in {path.name}")

    cost_name_row = header_row - 3 if header_row > 3 else None

    header_to_col: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        raw = ws.cell(header_row, c).value
        if raw is None:
            continue
        header_to_col[str(raw).strip()] = c

    cost_blocks: list[tuple[str, int, int]] = []
    c = 1
    while c <= ws.max_column:
        h = str(ws.cell(header_row, c).value or "").strip()
        if h == CURRENCY_COL and c + 1 <= ws.max_column:
            flat_h = str(ws.cell(header_row, c + 1).value or "").strip()
            if flat_h == FLAT_COL:
                name = ""
                if cost_name_row:
                    name = str(ws.cell(cost_name_row, c).value or "").strip()
                if not name:
                    for up in range(cost_name_row or header_row - 1, 0, -1):
                        cand = ws.cell(up, c).value
                        if cand:
                            name = str(cand).strip()
                            break
                cost_blocks.append((name or f"Cost@{c}", c, c + 1))
                c += 2
                continue
        c += 1

    lanes: list[dict] = []
    for r in range(header_row + 1, ws.max_row + 1):
        rate_col = header_to_col.get("RateCode")
        if not rate_col:
            break
        rate_code = ws.cell(r, rate_col).value
        if rate_code is None or str(rate_code).strip() == "":
            if all(
                ws.cell(r, col).value in (None, "")
                for col in range(1, min(10, ws.max_column + 1))
            ):
                continue
        details = {}
        for ra_header, our_col in RA_DETAIL_HEADER_MAP.items():
            col = header_to_col.get(ra_header)
            if not col:
                continue
            val = ws.cell(r, col).value
            details[our_col] = "" if val is None else val

        costs = {}
        for name, cur_c, flat_c in cost_blocks:
            cur = ws.cell(r, cur_c).value
            flat = ws.cell(r, flat_c).value
            costs[name] = {
                CURRENCY_COL: "" if cur is None else cur,
                FLAT_COL: "" if flat is None else flat,
            }

        lanes.append(
            {
                "RateCode": "" if rate_code is None else str(rate_code).strip(),
                "details": details,
                "costs": costs,
            }
        )

    wb.close()
    return lanes


def reorder_matrix_by_ra(
    matrix: pd.DataFrame,
    ra_lanes: list[dict],
    transport_cost_name: str,
) -> pd.DataFrame:
    """
    Sort matrix lanes like previous RA Rate card; append new lanes at the end.

    When several lanes share the same RateCode (e.g. Single vs Round trip),
    pair each RA occurrence to the best-matching matrix row by content.
    """
    attrs = dict(matrix.attrs)
    n = len(matrix)
    by_code: dict[str, list[int]] = defaultdict(list)
    for idx in range(n):
        code = (
            str(matrix.at[idx, "RateCode"]).strip()
            if pd.notna(matrix.at[idx, "RateCode"])
            else ""
        )
        by_code[code].append(idx)

    new_order: list[int] = []
    matched_ra: list[dict | None] = []
    used: set[int] = set()
    for lane in ra_lanes:
        code = lane["RateCode"]
        candidates = [i for i in by_code.get(code, []) if i not in used]
        idx = pick_best_lane_index(candidates, lane, matrix, transport_cost_name)
        if idx is None:
            continue
        new_order.append(idx)
        matched_ra.append(lane)
        used.add(idx)

    for idx in range(n):
        if idx not in used:
            new_order.append(idx)
            matched_ra.append(None)

    reordered = matrix.iloc[new_order].copy().reset_index(drop=True)
    reordered["lane #"] = range(1, len(reordered) + 1)

    for attr in PER_ROW_ATTRS:
        values = attrs.get(attr)
        if isinstance(values, list) and len(values) == n:
            attrs[attr] = [values[i] for i in new_order]

    attrs[ATTR_RA_MATCH] = matched_ra
    reordered.attrs.update(attrs)
    return reordered


def mark_changes_vs_ra(
    matrix: pd.DataFrame,
    ra_lanes: list[dict],
    transport_cost_name: str,
) -> pd.DataFrame:
    """
    Compare to previous RA:
      - new lanes (no paired RA occurrence) -> ATTR_NEW_LANE (yellow)
      - existing lanes with field diffs -> ATTR_CHANGED_CELLS (green)
    """
    cost_names: list[str] = list(matrix.attrs.get("cost_names", []))
    changed_rows: list[set[str]] = []
    new_flags: list[bool] = []

    paired: list[dict | None] | None = matrix.attrs.get(ATTR_RA_MATCH)
    if not isinstance(paired, list) or len(paired) != len(matrix):
        by_code: dict[str, list[dict]] = defaultdict(list)
        for lane in ra_lanes:
            by_code[lane["RateCode"]].append(lane)
        paired = []
        for _, row in matrix.iterrows():
            code = str(row["RateCode"]).strip() if pd.notna(row["RateCode"]) else ""
            candidates = by_code.get(code, [])
            if not candidates:
                paired.append(None)
                continue
            best_i = 0
            best_score = lane_match_score(
                row, candidates[0], matrix, transport_cost_name
            )
            for i, cand in enumerate(candidates[1:], start=1):
                score = lane_match_score(row, cand, matrix, transport_cost_name)
                if score > best_score:
                    best_i = i
                    best_score = score
            paired.append(candidates.pop(best_i))

    for r_idx, (_, row) in enumerate(matrix.iterrows()):
        prev = paired[r_idx] if r_idx < len(paired) else None
        changed: set[str] = set()
        if prev is None:
            new_flags.append(True)
            changed_rows.append(set())
            continue

        new_flags.append(False)
        for col in DETAIL_COLUMNS:
            if col == "lane #":
                continue
            old = prev["details"].get(col, "")
            if cmp_value(row[col]) != cmp_value(old):
                changed.add(col)

        for cost_name in cost_names:
            ra_vals = None
            for ra_name, vals in prev["costs"].items():
                if cost_names_match(cost_name, ra_name):
                    ra_vals = vals
                    break
            cur_key = currency_key(cost_name)
            f_key = flat_key(cost_name)
            our_cur = row[cur_key] if cur_key in matrix.columns else ""
            our_flat = row[f_key] if f_key in matrix.columns else ""
            if ra_vals is None:
                if cmp_value(our_cur) != "" or cmp_value(our_flat) != "":
                    if cmp_value(our_cur) != "":
                        changed.add(cur_key)
                    if cmp_value(our_flat) != "":
                        changed.add(f_key)
                continue
            if cmp_value(our_cur) != cmp_value(ra_vals.get(CURRENCY_COL, "")):
                changed.add(cur_key)
            if cmp_value(our_flat) != cmp_value(ra_vals.get(FLAT_COL, "")):
                changed.add(f_key)

        changed_rows.append(changed)

    matrix.attrs[ATTR_CHANGED_CELLS] = changed_rows
    matrix.attrs[ATTR_NEW_LANE] = new_flags
    return matrix


def apply_previous_ra_order_and_diff(
    sheets: dict[str, pd.DataFrame],
    transport_cost_name: str,
    ra_path: Path | None = None,
    hero_ra_path: Path | None = None,
    mark_duplicates=None,
    collect_generated_conditions=None,
) -> dict[str, pd.DataFrame]:
    """
    Sort/compare Matrix + Other against main previous RA,
    and HERO Group against its own previous RA.

    mark_duplicates / collect_generated_conditions are injected from the flow
    so common RA logic stays reusable.
    """
    ra_lanes = None
    if ra_path is not None:
        ra_lanes = load_previous_ra_rate_card(ra_path)
        print(f"Previous RA (Matrix) Rate card lanes: {len(ra_lanes)}")
    else:
        print("Previous RA (Matrix): skipped")

    hero_ra_lanes = None
    if hero_ra_path is not None:
        hero_ra_lanes = load_previous_ra_rate_card(hero_ra_path)
        print(f"Previous RA (HERO) Rate card lanes: {len(hero_ra_lanes)}")
    else:
        print("Previous RA (HERO): skipped")

    sheet_ra: dict[str, list[dict] | None] = {
        SHEET_MATRIX: ra_lanes,
        SHEET_OTHER: ra_lanes,
        SHEET_HERO: hero_ra_lanes,
    }

    for name, lanes in sheet_ra.items():
        if name not in sheets:
            continue
        matrix = sheets[name]
        if lanes is not None:
            matrix = reorder_matrix_by_ra(matrix, lanes, transport_cost_name)
            matrix = mark_changes_vs_ra(matrix, lanes, transport_cost_name)
            n_changed_rows = sum(
                1 for s in matrix.attrs.get(ATTR_CHANGED_CELLS, []) if s
            )
            n_new_rows = sum(1 for f in matrix.attrs.get(ATTR_NEW_LANE, []) if f)
            print(
                f"  [{name}] reordered vs previous RA; "
                f"changed rows: {n_changed_rows}; new rows: {n_new_rows}"
            )
        else:
            matrix.attrs[ATTR_CHANGED_CELLS] = [set() for _ in range(len(matrix))]
            matrix.attrs[ATTR_NEW_LANE] = [False] * len(matrix)
            matrix.attrs[ATTR_RA_MATCH] = [None] * len(matrix)
            print(f"  [{name}] previous RA skipped (no sort / no green/yellow highlight)")

        if mark_duplicates is not None:
            matrix = mark_duplicates(matrix)
            n_dup_rows = sum(1 for f in matrix.attrs.get(ATTR_DUPLICATE_LANE, []) if f)
            print(f"  [{name}] duplicate lanes: {n_dup_rows}")
        sheets[name] = matrix

    if SHEET_CONDITIONS in sheets and collect_generated_conditions is not None:
        frames = []
        if SHEET_MATRIX in sheets:
            frames.append(sheets[SHEET_MATRIX][DETAIL_COLUMNS])
        if SHEET_HERO in sheets:
            frames.append(sheets[SHEET_HERO][DETAIL_COLUMNS])
        company_conds = collect_generated_conditions(*frames)
        from flows.ftl.schema import ATTR_COMPANY_CODE_CONDITIONS

        sheets[SHEET_CONDITIONS].attrs[ATTR_COMPANY_CODE_CONDITIONS] = company_conds

    return sheets
