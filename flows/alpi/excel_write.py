"""ALPI Excel writer — detail columns + Currency / weight-bracket cost headers."""

from __future__ import annotations

from pathlib import Path

import pandas as pd

from flows.alpi.config import (
    ATTR_DEST_COMBINED,
    ATTR_DUPLICATE_LANE,
    CURRENCY_COL,
    DETAIL_COLUMNS,
    FLAT_COL,
    RATE_BY_LABEL,
    SHEET_MATRIX,
)
from flows.alpi.cost_utils import bracket_key, currency_key, flat_key
from flows.ftl import paths as project_paths
from flows.ftl.schema import (
    ATTR_CHANGED_CELLS,
    ATTR_NEW_LANE,
    ATTR_REMOVED_LANE,
    ATTR_VALID_FROM_RATES,
)


def write_matrix_sheet(ws, matrix: pd.DataFrame) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    cost_blocks: list[dict] = list(matrix.attrs.get("cost_blocks", []))
    duplicate_flags: list[bool] = list(matrix.attrs.get(ATTR_DUPLICATE_LANE, []))
    dest_combined_flags: list[bool] = list(matrix.attrs.get(ATTR_DEST_COMBINED, []))
    changed_rows: list[set] = list(matrix.attrs.get(ATTR_CHANGED_CELLS, []))
    new_flags: list[bool] = list(matrix.attrs.get(ATTR_NEW_LANE, []))
    removed_flags: list[bool] = list(matrix.attrs.get(ATTR_REMOVED_LANE, []))
    valid_from_rates_flags: list[bool] = list(
        matrix.attrs.get(ATTR_VALID_FROM_RATES, [])
    )

    header_rows = 3
    detail_count = len(DETAIL_COLUMNS)

    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    detail_fill = PatternFill("solid", fgColor="F2F2F2")
    cost_fill = PatternFill("solid", fgColor="D9D2E9")
    dup_fill = PatternFill("solid", fgColor="FFC7CE")
    dest_fill = PatternFill("solid", fgColor="BDD7EE")  # blue for gathered destinations
    change_fill = PatternFill("solid", fgColor="C6EFCE")
    new_fill = PatternFill("solid", fgColor="FFEB9C")
    removed_fill = PatternFill("solid", fgColor="F4B183")
    valid_from_fill = PatternFill("solid", fgColor="E2D5F1")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    dup_font = Font(color="9C0006")
    dest_font = Font(color="1F4E79", underline="single")
    change_font = Font(color="006100")
    new_font = Font(color="9C5700")
    removed_font = Font(color="833C0C")
    valid_from_font = Font(color="7030A0")

    for col_idx, title in enumerate(DETAIL_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.fill = detail_fill
        cell.alignment = center
        cell.font = Font(bold=True)
        cell.border = border
        ws.merge_cells(
            start_row=1, start_column=col_idx, end_row=header_rows, end_column=col_idx
        )
        for row in range(1, header_rows + 1):
            ws.cell(row=row, column=col_idx).border = border
            ws.cell(row=row, column=col_idx).fill = detail_fill

    col_idx = detail_count + 1
    for block in cost_blocks:
        name = block["name"]
        if block["kind"] == "flat":
            span = 2
            sub_headers = [CURRENCY_COL, FLAT_COL]
            data_keys = [currency_key(name), flat_key(name)]
        else:
            brackets = list(block.get("brackets", []))
            span = 1 + len(brackets)
            sub_headers = [CURRENCY_COL] + brackets
            data_keys = [currency_key(name)] + [
                bracket_key(name, br) for br in brackets
            ]

        left = col_idx
        right = col_idx + span - 1

        name_cell = ws.cell(row=1, column=left, value=name)
        name_cell.fill = cost_fill
        name_cell.alignment = center
        name_cell.font = Font(bold=True)
        if span > 1:
            ws.merge_cells(
                start_row=1, start_column=left, end_row=1, end_column=right
            )

        rate_cell = ws.cell(row=2, column=left, value=RATE_BY_LABEL)
        rate_cell.fill = cost_fill
        rate_cell.alignment = center
        if span > 1:
            ws.merge_cells(
                start_row=2, start_column=left, end_row=2, end_column=right
            )

        for offset, sub in enumerate(sub_headers):
            cell = ws.cell(row=3, column=left + offset, value=sub)
            cell.fill = cost_fill
            cell.alignment = center
            cell.font = Font(bold=True)

        for row in range(1, header_rows + 1):
            for c in range(left, right + 1):
                ws.cell(row=row, column=c).border = border
                ws.cell(row=row, column=c).fill = cost_fill

        block["_data_keys"] = data_keys
        col_idx = right + 1

    for r_idx, (_, row) in enumerate(matrix.iterrows()):
        excel_row = header_rows + 1 + r_idx
        is_duplicate = r_idx < len(duplicate_flags) and duplicate_flags[r_idx]
        is_dest_combined = (
            r_idx < len(dest_combined_flags) and dest_combined_flags[r_idx]
        )
        is_new = r_idx < len(new_flags) and new_flags[r_idx]
        is_removed = r_idx < len(removed_flags) and removed_flags[r_idx]
        changed = changed_rows[r_idx] if r_idx < len(changed_rows) else set()

        for c_idx, col_name in enumerate(DETAIL_COLUMNS, start=1):
            value = row[col_name]
            cell = ws.cell(
                row=excel_row,
                column=c_idx,
                value="" if pd.isna(value) else value,
            )
            cell.border = border
            if is_removed:
                cell.fill = removed_fill
                cell.font = removed_font
            elif (
                col_name == "Valid from"
                and r_idx < len(valid_from_rates_flags)
                and valid_from_rates_flags[r_idx]
            ):
                cell.fill = valid_from_fill
                cell.font = valid_from_font
            elif (not is_new) and col_name in changed:
                cell.fill = change_fill
                cell.font = change_font
            elif is_new:
                cell.fill = new_fill
                cell.font = new_font
            elif is_duplicate:
                cell.fill = dup_fill
                cell.font = dup_font
            elif is_dest_combined and col_name == "DestinationNo":
                cell.fill = dest_fill
                cell.font = dest_font

        col_idx = detail_count + 1
        for block in cost_blocks:
            for key in block.get("_data_keys", []):
                value = row[key] if key in matrix.columns else ""
                cell = ws.cell(
                    row=excel_row,
                    column=col_idx,
                    value="" if pd.isna(value) else value,
                )
                cell.border = border
                if is_removed:
                    cell.fill = removed_fill
                    cell.font = removed_font
                elif (not is_new) and key in changed:
                    cell.fill = change_fill
                    cell.font = change_font
                elif is_new:
                    cell.fill = new_fill
                    cell.font = new_font
                elif is_duplicate:
                    cell.fill = dup_fill
                    cell.font = dup_font
                col_idx += 1

    total_cols = col_idx - 1
    for c_idx in range(1, detail_count + 1):
        ws.column_dimensions[get_column_letter(c_idx)].width = 14
    for c_idx in range(detail_count + 1, total_cols + 1):
        ws.column_dimensions[get_column_letter(c_idx)].width = 10

    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 18
    if total_cols >= 1:
        ws.auto_filter.ref = (
            f"A{header_rows}:{get_column_letter(max(total_cols, 1))}{header_rows}"
        )
    ws.freeze_panes = f"A{header_rows + 1}"


def save_workbook(sheets: dict[str, pd.DataFrame], rates_file: Path) -> Path:
    from openpyxl import Workbook

    project_paths.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out_path = project_paths.OUTPUT_DIR / f"{rates_file.stem}_matrix.xlsx"

    wb = Workbook()
    default = wb.active
    wb.remove(default)

    for sheet_name, matrix in sheets.items():
        ws = wb.create_sheet(sheet_name)
        write_matrix_sheet(ws, matrix)

    wb.save(out_path)
    return out_path
